"""Branding (and every other former Company Setting) is now static —
backend/branding.py — there is no Settings page or table any more. This
proves every surface that renders it (a schedule PDF, a payslip PDF, a
report PDF, an Excel export, the company block sent to the email
service) shows the fixed brand.
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import branding

OLD_HARDCODED_NAME = "Microfinance LMS"


@pytest.fixture()
def staff(make_user, auth_headers):
    people = {"ceo": make_user("ceo"), "maker": make_user("maker"), "checker": make_user("checker")}
    users = {role: pair[0] for role, pair in people.items()}
    headers = {role: auth_headers(pair[0].email, pair[1]) for role, pair in people.items()}
    return users, headers


def pdf_text(stream):
    reader = PdfReader(stream)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


# --------------------------------------------------------------- PDFs


def test_schedule_pdf_shows_the_static_brand(client, staff):
    _, headers = staff
    borrower = client.post("/api/borrowers", json={
        "name": "Branding Borrower", "phone": "0711000111", "id_number": "BR-1",
    }, headers=headers["maker"])
    assert borrower.status_code == 201, borrower.get_json()
    loan = client.post("/api/loans", json={
        "borrower_id": borrower.get_json()["borrower"]["id"],
        "principal_amount": 500000, "interest_rate": 10, "term_months": 6, "start_date": "2026-01-01",
    }, headers=headers["maker"])
    assert loan.status_code == 201, loan.get_json()
    loan_id = loan.get_json()["loan"]["id"]

    response = client.get(f"/api/loans/{loan_id}/schedule/export", headers=headers["maker"])
    assert response.status_code == 200
    text = pdf_text(BytesIO(response.data))
    assert branding.COMPANY_NAME in text
    assert OLD_HARDCODED_NAME not in text


def test_payslip_pdf_shows_the_static_brand(client, staff):
    _, headers = staff
    employee = client.post("/api/employees", json={
        "name": "Payslip Branding Employee", "phone": "0700222333", "job_title": "Officer",
        "salary": 500000, "start_date": "2026-01-01",
    }, headers=headers["ceo"])
    assert employee.status_code == 201, employee.get_json()
    batch = client.post("/api/payroll", json={"month": "2026-05"}, headers=headers["ceo"])
    assert batch.status_code == 201, batch.get_json()
    batch_id = batch.get_json()["payroll_batch"]["id"]
    line_id = batch.get_json()["payroll_batch"]["lines"][0]["id"]
    finalized = client.post(f"/api/payroll/{batch_id}/finalize", headers=headers["ceo"])
    assert finalized.status_code == 200, finalized.get_json()

    response = client.get(f"/api/payroll/lines/{line_id}/payslip", headers=headers["ceo"])
    assert response.status_code == 200
    text = pdf_text(BytesIO(response.data))
    assert branding.COMPANY_NAME in text


def test_report_pdf_shows_the_static_brand(client, staff):
    _, headers = staff
    response = client.get("/api/reports/portfolio-at-risk?format=pdf", headers=headers["ceo"])
    assert response.status_code == 200
    text = pdf_text(BytesIO(response.data))
    assert branding.COMPANY_NAME in text


# --------------------------------------------------------------- Excel


def test_report_excel_shows_the_static_brand(client, staff):
    _, headers = staff
    response = client.get("/api/reports/portfolio-at-risk?format=excel", headers=headers["ceo"])
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    summary = workbook["Summary"]
    cell_values = [cell.value for row in summary.iter_rows() for cell in row if cell.value]
    assert any(branding.COMPANY_NAME == v for v in cell_values)
    header_fill = workbook[workbook.sheetnames[1]]["A1"].fill.fgColor.rgb
    assert branding.PRIMARY_COLOR.lstrip("#").upper() in str(header_fill).upper()


# --------------------------------------------------------------- email company block


def test_email_company_block_always_carries_the_static_brand(client, staff, monkeypatch, app):
    captured = {}

    class FakeResponse:
        status_code = 200

        def json(self):
            return {"status": "sent"}

    def fake_post(url, headers=None, json=None, timeout=None):
        captured["payload"] = json
        return FakeResponse()

    monkeypatch.setattr("services.email_client.requests.post", fake_post)

    with app.app_context():
        from services.email_client import send_email
        send_email("password_reset", "someone@example.com", {"name": "X", "reset_link": "http://x"})

    company = captured["payload"]["company"]
    assert company["name"] == branding.COMPANY_NAME
    assert company["primary_color"] == branding.PRIMARY_COLOR
    assert company["accent_color"] == branding.ACCENT_COLOR
    assert company["footer_text"] == branding.FOOTER_TEXT
    assert "logo_base64" in company


def test_email_company_block_is_always_built_from_static_branding(client, monkeypatch, app):
    captured = {}

    class FakeResponse:
        status_code = 200

        def json(self):
            return {"status": "sent"}

    def fake_post(url, headers=None, json=None, timeout=None):
        captured["payload"] = json
        return FakeResponse()

    monkeypatch.setattr("services.email_client.requests.post", fake_post)

    with app.app_context():
        from services.email_client import send_email
        result = send_email("password_reset", "someone@example.com", {"name": "X", "reset_link": "http://x"})

    assert result == {"status": "sent"}
    assert captured["payload"]["company"]["name"] == branding.COMPANY_NAME
